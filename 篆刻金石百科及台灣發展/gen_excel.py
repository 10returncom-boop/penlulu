# -*- coding: utf-8 -*-
"""
篆刻金石百科｜台灣發展 — gen_excel.py
================================================
產出兩份 Excel 交付物（與 build 完全同步，讀 data/site.json）：
  1) 網站地圖.xlsx           全部頁面清單（網址/主題/子分類/標題/標籤/難度/更新）
  2) 網站變數與網站導覽.xlsx  全站變數＋導覽樹（章節→子分類→條目）＋標籤索引
"""
import os
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
site = json.load(open(os.path.join(ROOT, 'data', 'site.json'), encoding='utf-8'))

HEAD_FILL = PatternFill('solid', fgColor='9B2C2C')
HEAD_FONT = Font(color='FFFFFF', bold=True, name='Microsoft JhengHei')
CELL_FONT = Font(name='Microsoft JhengHei', size=10)
THIN = Side(style='thin', color='DDDDDD')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'


def autofit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fill_rows(ws, rows):
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical='center', wrap_text=True)


# ============ 1) 網站地圖.xlsx ============
wb1 = Workbook()
ws = wb1.active
ws.title = '網站地圖'
ws.append(['頁面網址', '層級', '大主題', '子分類', '標題', '標籤', '難度', '更新日期'])
style_header(ws, 8)

rows = []
rows.append(['index.html', '首頁', '—', '—', site['site']['name'], '—', '—', site['site']['updated']])
rows.append(['search.html', '工具頁', '—', '—', '全站搜尋＋多面向篩選', '—', '—', site['site']['updated']])
rows.append(['az-index.html', '工具頁', '—', '—', 'A–Z 字母索引', '—', '—', site['site']['updated']])
rows.append(['tags/index.html', '工具頁', '—', '—', '標籤彙整', '—', '—', site['site']['updated']])
rows.append(['404.html', '工具頁', '—', '—', '404 找不到頁面', '—', '—', site['site']['updated']])

chmap = {c['id']: c for c in site['chapters']}
submap = {}
for c in site['chapters']:
    for s in c['subs']:
        submap[(c['id'], s['id'])] = s

for c in site['chapters']:
    rows.append([c['id'] + '/index.html', '章節頁', c['zh'], '—', c['zh'], '—', '—', site['site']['updated']])
    for s in c['subs']:
        rows.append([c['id'] + '/' + s['id'] + '/index.html', '子分類頁', c['zh'], s['zh'], s['zh'], '—', '—', site['site']['updated']])

for a in site['articles']:
    c = chmap[a['chapter']]
    s = submap[(a['chapter'], a['sub'])]
    rows.append([a['chapter'] + '/' + a['sub'] + '/' + a['slug'] + '.html', '文章頁',
                 c['zh'], s['zh'], a['title'], a['tags'], a['difficulty'], site['site']['updated']])

fill_rows(ws, rows)
autofit(ws, [42, 10, 16, 16, 26, 24, 8, 12])
wb1.save(os.path.join(ROOT, '網站地圖.xlsx'))
print('✓ 網站地圖.xlsx（%d 列）' % len(rows))

# ============ 2) 網站變數與網站導覽.xlsx ============
wb2 = Workbook()

# Sheet A：全站變數
ws1 = wb2.active
ws1.title = '全站變數'
ws1.append(['變數名稱', '值', '說明'])
style_header(ws1, 3)
var_rows = [
    ['site_name', site['site']['name'], '網站名稱'],
    ['site_name_en', site['site']['name_en'], '英文名稱'],
    ['site_desc', site['site']['desc'], '網站描述（description）'],
    ['root_url', site['site']['rootUrl'], '網站網域（canonical）'],
    ['version', site['site']['version'], '版本'],
    ['updated', site['site']['updated'], '建置日期'],
    ['chapter_count', str(len(site['chapters'])), '大主題數'],
    ['article_count', str(len(site['articles'])), '文章數'],
    ['theme_default', 'light（白天模式）', '預設主題'],
    ['generator', 'build.py（Python 零依賴靜態站）', '產生器'],
]
fill_rows(ws1, var_rows)
autofit(ws1, [18, 44, 30])

# Sheet B：導覽樹
ws2 = wb2.create_sheet('導覽樹')
ws2.append(['大主題', '子分類', '條目標題', '條目網址', '難度'])
style_header(ws2, 5)
nav_rows = []
for c in site['chapters']:
    nav_rows.append([c['zh'], '—', '（章節首頁）', c['id'] + '/index.html', '—'])
    for s in c['subs']:
        nav_rows.append(['', s['zh'], '（子分類首頁）', c['id'] + '/' + s['id'] + '/index.html', '—'])
        for a in site['articles']:
            if a['chapter'] == c['id'] and a['sub'] == s['id']:
                nav_rows.append(['', '', a['title'], a['chapter'] + '/' + a['sub'] + '/' + a['slug'] + '.html', a['difficulty']])
fill_rows(ws2, nav_rows)
autofit(ws2, [16, 16, 28, 44, 8])

# Sheet C：標籤索引
ws3 = wb2.create_sheet('標籤索引')
ws3.append(['標籤', '文章數', '範例文章'])
style_header(ws3, 3)
tag_map = {}
for a in site['articles']:
    for t in [x.strip() for x in a['tags'].split(',') if x.strip()]:
        tag_map.setdefault(t, []).append(a['title'])
tag_rows = [[t, len(v), '、'.join(v[:3])] for t, v in sorted(tag_map.items(), key=lambda x: -len(x[1]))]
fill_rows(ws3, tag_rows)
autofit(ws3, [18, 8, 50])

wb2.save(os.path.join(ROOT, '網站變數與網站導覽.xlsx'))
print('✓ 網站變數與網站導覽.xlsx（3 個工作表）')
