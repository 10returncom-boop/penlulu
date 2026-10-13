# -*- coding: utf-8 -*-
"""
產出 規格書.docx — 篆刻金石百科｜台灣發展 完整規格書。
全部數據（站名、章節、文章數、圖片數、URL 數）動態讀取，改資料不必改本檔。
"""
import os
import json
import glob
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.abspath(__file__))
site = json.load(open(os.path.join(ROOT, 'data', 'site.json'), encoding='utf-8'))

# 實際統計（與建置同步）
html_files = glob.glob(os.path.join(ROOT, '**', '*.html'), recursive=True)
webp_files = glob.glob(os.path.join(ROOT, 'assets', 'images', '*.webp'))
md_files = glob.glob(os.path.join(ROOT, 'data', 'articles', '**', '*.md'), recursive=True)
site_urls = len(site['articles']) + len(site['chapters']) + sum(len(c['subs']) for c in site['chapters']) + 5

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Microsoft JhengHei'
style.font.size = Pt(11)


def h1(t):
    return doc.add_heading(t, level=1)


def h2(t):
    return doc.add_heading(t, level=2)


def para(t):
    return doc.add_paragraph(t)


# 封面
title = doc.add_heading(site['site']['name'], level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph('專案規格書（SPEC）')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta = doc.add_paragraph('版本 %s ｜ 建置日期 %s ｜ 零依賴靜態站生成器（Content Collections 式資料分離）' % (
    site['site']['version'], site['site']['updated']))
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

h1('一、專案定位')
para('網站名稱：%s（%s）' % (site['site']['name'], site['site']['name_en']))
para('一句話定位：%s' % site['site']['desc'])
para('目標：打造一個權威、專業、大白話、資訊齊全、差異化、圖文並茂的篆刻金石＋台灣發展知識網站。')
para('技術架構：Python 零依賴靜態生成器（build.py + data/ + CONTENT-SPEC.md），資料與模板分離，'
     '改資料不用動模板；data/ 即 Content Collections 內容層，未來可無痛升級 Astro / Next.js。')

h1('二、檔案與資料夾架構')
para('一個大主題＝一個子資料夾；一次點擊＝一頁專業分析。')
arch = [
    'build.py          建置腳本（讀 data/site.json + md，產生全站 HTML、sitemap.xml、config.js、manifest）',
    'seed.py           內容展開器（articles_data.py → site.json + 逐篇 md；台灣章節自動加入在地段落）',
    'data_master.py    %d 大主題／%d 子分類定義（唯一內容來源之一）' % (len(site['chapters']), sum(len(c['subs']) for c in site['chapters'])),
    'articles_data.py  %d 篇條目明細（唯一內容來源之二）' % len(site['articles']),
    'gen_images.py     PIL 產生全部 WebP（章節主視覺＋每篇主圖＋辨識卡）',
    'gen_excel.py      產出「網站地圖.xlsx」「網站變數與網站導覽.xlsx」',
    'gen_spec_doc.py   產出本檔 規格書.docx',
    'check.py          三項批次檢查（SEO／連結與圖片路徑／本機斷鏈）',
    'data/             site.json + articles/**/*.md（Content Collections 內容層）',
    'assets/css|js|images/  樣式、互動腳本、WebP 圖片（%d 張）' % len(webp_files),
    '網站地圖.xlsx      全部頁面清單',
    '網站變數與網站導覽.xlsx  變數＋導覽樹＋標籤索引',
    'CONTENT-SPEC.md   內容撰寫規範',
    'README.md         建置說明',
]
for a in arch:
    doc.add_paragraph(a, style='List Bullet')

h1('三、資訊架構（%d 大主題）' % len(site['chapters']))
for c in site['chapters']:
    h2('%s（%s）' % (c['zh'], c['en']))
    para(c['desc'])
    for s in c['subs']:
        n = sum(1 for a in site['articles'] if a['chapter'] == c['id'] and a['sub'] == s['id'])
        doc.add_paragraph('%s：%s（%d 篇）' % (s['zh'], s['desc'], n), style='List Bullet')

h1('四、UI 元件與功能清單')
feats = [
    'Breadcrumb 麵包屑：搭配 BreadcrumbList JSON-LD，搜尋結果可顯示路徑。',
    '內頁卡片：Accordion 折疊卡片。',
    '側邊元件：Sidebar 樹狀導覽 + TOC 文章目錄（Treemap 樹狀地圖）。',
    '主題樣式：預設白天模式（light），可切換暗色／跟隨系統。',
    'MegaMenu 巨型導覽：3 層以內簡潔下拉選單。',
    'Tag System 標籤系統：tags/ 頁 + 標籤雲。',
    '全站搜尋：支援自動完成 autocomplete。',
    'Faceted Search 多面向篩選：主題／難度／關鍵字。',
    '相關文章推薦模組。',
    '知識圖譜／收藏清單功能（localStorage）。',
    '右下角固定站內導航：回首頁、站內導覽、回到頂端、分頁 pager。',
    'OG 標籤、Favicon 完整設定（OG=9）。',
]
for f in feats:
    doc.add_paragraph(f, style='List Bullet')

h1('五、結構化資料與 SEO')
seo = [
    '每頁含 ≥3 種 JSON-LD：BreadcrumbList + Article/CollectionPage + FAQPage/WebSite/ItemList。',
    'FAQPage JSON-LD：每篇 4 組 Q/A 自動結構化。',
    '完整 sitemap.xml（%d 條 URL）／robots.txt／404 頁／webmanifest。' % site_urls,
    '頁面標題 H1/H2/H3 階層，大白話、去 AI 模板化、台灣在地觀點。',
    '圖片統一 WebP（%d 張），全數附 alt text。' % len(webp_files),
    '載入最佳化（WebP＋lazy loading）、AEO/RAG 檢索友好、Topic-Cluster 主題群聚。',
]
for s in seo:
    doc.add_paragraph(s, style='List Bullet')

h1('六、自動檢查腳本（三項檢核）')
chk = [
    '掃描全部 HTML SEO 元素：每頁 title、meta description、OG、JSON-LD。',
    '檢查所有內部／外部連結、圖片檔案路徑。',
    '比對本機檔案存在與否，自動偵測、列出 404 斷鏈與遺失圖片。',
]
for c in chk:
    doc.add_paragraph(c, style='List Bullet')
para('執行方式：python check.py（或 python check.py --summary），目標為 0 SEO 缺失、0 斷鏈、0 遺失圖片。')

h1('七、批次更新機制')
upd = [
    '改資料：只動 data_master.py（章節）與 articles_data.py（條目），不動模板。',
    '重建：python seed.py → python gen_images.py → python build.py → python gen_excel.py → python gen_spec_doc.py',
    '檢核：python check.py（批次掃描全部 HTML、連結與本機檔案）。',
    '可重複執行：N 頁一次建置，JSON（site.json）與 Excel（網站地圖）同步更新。',
]
for u in upd:
    doc.add_paragraph(u, style='List Bullet')

h1('八、交付物清單')
deliv = [
    '完整多頁網站原始程式（Python + Content Collections 式資料分離，%d 個 HTML 頁）' % len(html_files),
    '規格書（本檔 規格書.docx）',
    '網站地圖.xlsx（全部頁面）',
    '網站變數與網站導覽.xlsx（變數＋導覽樹＋標籤）',
    'CONTENT-SPEC.md 內容規範、README.md 建置說明',
    '批次更新機制：可重複執行內容重建與斷鏈檢查',
]
for d in deliv:
    doc.add_paragraph(d, style='List Bullet')

out = os.path.join(ROOT, '規格書.docx')
doc.save(out)
print('✓ 規格書.docx 已產出：%s（%d 章、%d 篇）' % (out, len(site['chapters']), len(site['articles'])))
